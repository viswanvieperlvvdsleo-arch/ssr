from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SJ_Info_Business_Solutions_Platform_User_Guide.docx"
LOGO = ROOT / "public" / "SJINFOBUSINESSSOLUTIONSLOGO.png"

BLUE = "0A6ED1"
BLUE_DARK = "124A7A"
GREEN = "0B8F55"
GOLD = "C88A00"
INK = "172033"
TEXT = "334155"
MUTED = "64748B"
LIGHT_BLUE = "EAF3FF"
LIGHT_GREEN = "EAF8F1"
LIGHT_GOLD = "FFF7DA"
LIGHT_GRAY = "F4F7FA"
GRID = "CDD7E3"
WHITE = "FFFFFF"
RED = "B42318"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, color=TEXT, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border(paragraph, side="bottom", color=GRID, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE_DARK, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, BLUE_DARK, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(3)
    r = header_p.add_run("SJ INFO BUSINESS SOLUTIONS")
    set_run_font(r, size=9, color=BLUE_DARK, bold=True)
    r = header_p.add_run("  |  PLATFORM USER GUIDE")
    set_run_font(r, size=9, color=MUTED, bold=True)
    set_paragraph_border(header_p, color=GRID, size=6, space=5)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(3)
    r = footer_p.add_run("Client guide  |  Page ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(footer_p)


def add_title(doc, text, size=30, color=INK, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, color=GREEN, bold=True)
    return p


def add_subtitle(doc, text, size=13.5, after=18):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, size=size, color=MUTED)
    return p


def add_text(doc, text, bold_prefix=None, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=INK)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_steps(doc, items):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id:08X}")
    abstract.append(nsid)
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)

    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        p_pr.append(num_pr)
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, label, text, kind="info"):
    fill, border, label_color = {
        "info": (LIGHT_BLUE, BLUE, BLUE_DARK),
        "success": (LIGHT_GREEN, GREEN, GREEN),
        "warning": (LIGHT_GOLD, GOLD, "7A5500"),
        "danger": ("FEECEB", RED, RED),
    }[kind]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, fill)
    set_paragraph_border(p, side="left", color=border, size=18, space=5)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=TEXT)
    return p


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=BLUE_DARK, bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "F8FAFC")
        for index, value in enumerate(row):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=TEXT, bold=(index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_section_title(doc, title, intro=None):
    p = doc.add_paragraph(title, style="Heading 1")
    set_paragraph_border(p, color=GRID, size=6, space=6)
    if intro:
        add_text(doc, intro, after=9)


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(20)
    p.add_run().add_picture(str(LOGO), width=Inches(2.25))
    add_kicker(doc, "Client operations guide")
    add_title(doc, "SJ INFO BUSINESS SOLUTIONS", size=31, color=INK, after=5)
    add_title(doc, "Platform User Guide", size=23, color=BLUE_DARK, after=10)
    add_subtitle(doc, "How to use every page, manage daily workflows, and understand live updates across the platform.", size=13.5, after=24)
    add_table(
        doc,
        ["Document", "Details"],
        [
            ("Prepared for", "SJ Info Business Solutions"),
            ("Audience", "Administrators, employees, trainers, and participants"),
            ("Edition", "Version 1.0 | September 2026"),
            ("Purpose", "Client handover, onboarding, and daily operations"),
        ],
        [2700, 6660],
        font_size=10,
    )
    add_callout(doc, "Scope", "This guide explains what users see and how actions update across devices. It intentionally does not describe source code, databases, or technical implementation.", "success")

    add_page_break(doc)

    # Guide map and roles
    add_section_title(doc, "1. Guide Map and Access", "The platform changes what it shows according to the signed-in account, assigned permissions, and meeting or training membership.")
    doc.add_paragraph("Who sees what", style="Heading 2")
    add_table(
        doc,
        ["Role", "Main access", "Typical responsibility"],
        [
            ("Super Admin", "All pages and all records", "Overall control, account setup, data oversight, reports"),
            ("Admin", "Administrative pages and operations", "Users, content, meetings, payments, requests, reports"),
            ("Employee", "Permission-based staff tools", "Posts, sourcing, services, meetings, requests, task work"),
            ("Trainer", "Allocated training and profile tools", "Join assigned meetings, attendance, trainer profile"),
            ("Participant / User", "Personal and allocated content", "Feed, chat, services, assigned meetings, own history"),
        ],
        [1500, 3300, 4560],
        font_size=9.2,
    )
    add_callout(doc, "Access rule", "A page or button that is not permitted for an account is hidden or unavailable. Employees may receive selected permissions instead of full administrative access.", "warning")
    doc.add_paragraph("Pages covered", style="Heading 2")
    add_bullets(doc, [
        "Access and installation, login, navigation, profile, and notifications",
        "Feed, internal requirements, Task Board, and employee reporting",
        "Chat, groups, message status, scheduled messages, and chat requests",
        "Services, server-access purchase flow, bookmarks, and payment history",
        "Meetings, meeting room controls, attendance, and training reports",
        "Trainers / Users, Account Management, Data Management, and Settings",
    ])
    doc.add_paragraph("Reading convention", style="Heading 2")
    add_text(doc, "Staff only: Actions marked as staff-only are intended for Super Admin, Admin, or authorized Employee accounts.", bold_prefix="Staff only:")
    add_text(doc, "Live: The screen updates automatically while connected. A short delay of a few seconds is normal between devices.", bold_prefix="Live:")
    add_text(doc, "Push: A device notification may appear even when the app is closed, provided notifications are allowed on that device.", bold_prefix="Push:")

    add_page_break(doc)

    # Access and navigation
    add_section_title(doc, "2. Access, Install, and Navigation", "Users can open the platform in a browser or install it as an app on a supported phone or computer.")
    doc.add_paragraph("Open or install", style="Heading 2")
    add_steps(doc, [
        "Open the company-provided platform link.",
        "Choose Open in browser for immediate access, or Download App to install it on the device.",
        "Allow notifications, camera, and microphone when the device asks. Screen-sharing permission is requested only when Share is used.",
        "Sign in with the email and password issued for the account.",
    ])
    add_callout(doc, "Best practice", "Install the app on frequently used devices and allow notifications. This gives the most consistent reminders, chat alerts, and meeting actions.", "info")
    doc.add_paragraph("Navigation", style="Heading 2")
    add_bullets(doc, [
        "Desktop: use the left navigation menu. Chat may remain visible beside the selected workspace.",
        "Mobile: use the bottom navigation bar. Swipe the bar horizontally when more staff pages are available.",
        "Bell: opens notifications and shows the unread count.",
        "Profile image: opens profile, settings, account controls, and logout options.",
        "Back controls: return to the previous list, report, chat, or meeting details without losing the current account session.",
    ])
    doc.add_paragraph("Sign-in behavior", style="Heading 2")
    add_text(doc, "The account remains signed in on the device until logout, account restriction, password change requiring re-authentication, or browser storage is cleared. On a shared device, always use Logout when finished.")
    add_callout(doc, "Security", "Never share an administrator password. Create a separate account for each person so activity, attendance, messages, and permissions remain attributable.", "warning")

    add_page_break(doc)

    # Feed
    add_section_title(doc, "3. Feed", "The Feed is the main publishing area for announcements, training updates, discussions, videos, and staff-only requirements.")
    doc.add_paragraph("Browse and interact", style="Heading 2")
    add_bullets(doc, [
        "Use All to view the combined feed, or choose a category to narrow the list.",
        "Open a post to read details, view media, add a comment, like, share, or save it.",
        "Use the bookmark action to add a post to Bookmarks for later review.",
        "Internal Feed appears only to staff. Participants and public users do not see the tab or its posts.",
    ])
    doc.add_paragraph("Create a post", style="Heading 2")
    add_steps(doc, [
        "Select Create Post.",
        "Choose the category and enter a title and description.",
        "Choose Public for all permitted accounts or Internal for staff-only visibility.",
        "Optionally attach an image or video.",
        "For an internal staffing requirement, enable Create as a requirement before publishing.",
        "Publish the post. It appears to the correct audience and updates other active devices automatically.",
    ])
    doc.add_paragraph("Requirement indicator", style="Heading 2")
    add_table(
        doc,
        ["Indicator", "Meaning", "What users can do"],
        [
            ("White", "Open requirement", "An employee can choose Work on this."),
            ("Yellow", "Sourcing in progress", "Shows the workers and their progress; more employees may join."),
            ("Green", "Requirement completed", "Closed record; it cannot return to white or yellow."),
        ],
        [1500, 2800, 5060],
    )
    add_callout(doc, "Live behavior", "Likes, comments, new posts, and requirement status changes propagate to connected users without a manual refresh. If a device was offline, the latest state appears when it reconnects.", "success")

    add_page_break(doc)

    # Task Board
    add_section_title(doc, "4. Task Board", "The Task Board turns internal requirement posts into accountable sourcing work. It is available to staff accounts only.")
    doc.add_paragraph("Work a requirement", style="Heading 2")
    add_steps(doc, [
        "Open Task Board and choose Active, Completed, or All.",
        "Open a requirement card to see the creator, team, status, and sourcing order.",
        "Choose Work on this from the feed or task card to join the sourcing queue.",
        "Add a candidate profile with the required text, optional file, and optional @mention.",
        "Use the left and right controls to select a worker. The profile list below changes to the selected worker, so each employee's work remains separate.",
        "Choose Complete my work when that employee has finished sourcing. Authorized staff can close the complete requirement permanently.",
    ])
    doc.add_paragraph("Candidate profile decisions", style="Heading 2")
    add_table(
        doc,
        ["Action", "Use it when", "Who is informed"],
        [
            ("Follow up", "The profile should move to client discussion or interview follow-up.", "Requirement creator, sourcing workers, relevant staff, and administrators"),
            ("Rejected", "The profile is not proceeding.", "The same working group and administrators"),
            ("Completed", "The profile successfully finishes its workflow.", "The same working group and administrators"),
            ("View file", "A document was attached and needs review.", "No status change; the selected file opens"),
        ],
        [1500, 3800, 4060],
        font_size=9.2,
    )
    add_callout(doc, "Notification", "Profile additions, mentions, follow-up, rejection, completion, worker changes, and requirement closure generate targeted alerts. Selecting View opens the relevant task and profile.", "info")

    add_page_break(doc)

    # Chat
    add_section_title(doc, "5. Chat", "Chat supports direct, group, and service conversations with live delivery state and media sharing.")
    doc.add_paragraph("Start and use a chat", style="Heading 2")
    add_bullets(doc, [
        "Search existing conversations by account or group name.",
        "Open a direct conversation, create a group when permitted, or use the service chat created for support and purchases.",
        "Send text, images, videos, audio, or documents. Upload progress appears for larger files.",
        "Reply to a specific message, react, edit your message, forward, or delete according to available permissions.",
        "Open group details to review members, admins, shared media, documents, links, and chat settings.",
        "Schedule a meeting from a group so the group is preselected as the audience.",
    ])
    doc.add_paragraph("Message status", style="Heading 2")
    add_table(
        doc,
        ["Symbol", "Meaning"],
        [
            (".", "Sent. The recipient is currently offline or has not become active in the app."),
            ("..", "Delivered to at least one online recipient."),
            ("...", "Seen by all intended recipients."),
            ("... while sending", "The message is still being uploaded or submitted."),
        ],
        [1500, 7860],
    )
    doc.add_paragraph("Scheduled messages", style="Heading 2")
    add_steps(doc, [
        "Open the scheduling control in the conversation.",
        "Enter the message, date, time, and optional repeat pattern.",
        "Confirm the schedule. The message is sent by the platform at the due time, even if the sender is not actively viewing the chat.",
        "Review or cancel a scheduled message before it is sent when the option is available.",
    ])
    add_callout(doc, "Live behavior", "Active chats normally synchronize within a few seconds. Presence, unread counts, delivery dots, and seen status may briefly lag when a device is backgrounded, loses internet, or is closing.", "success")

    add_page_break(doc)

    # Notifications
    add_section_title(doc, "6. Notifications", "Notifications keep users informed about messages, meetings, posts, tasks, mentions, and profile decisions.")
    doc.add_paragraph("Notification Center", style="Heading 2")
    add_bullets(doc, [
        "Select the bell to open the notification list and unread count.",
        "Select an alert to open the exact chat, post, meeting, task, or candidate profile.",
        "Use Mark all read to clear the unread state without deleting records.",
        "Use delete controls to remove notification entries from the account's notification list.",
    ])
    doc.add_paragraph("Device notification actions", style="Heading 2")
    add_table(
        doc,
        ["Alert type", "Available actions"],
        [
            ("Chat", "Reply or Mark as read"),
            ("Meeting", "Start or Cancel / dismiss the reminder"),
            ("Post", "Like or View post"),
            ("Task / profile", "View the related task or profile"),
        ],
        [2400, 6960],
    )
    add_callout(doc, "Device requirement", "Push alerts require browser or app notification permission. Sound and vibration also depend on the phone's notification channel, volume, Do Not Disturb, and battery settings.", "warning")
    add_text(doc, "If Reply does not open a text field on a particular device, open the alert and reply inside Chat. Notification action support varies by operating system and browser.")

    add_page_break(doc)

    # Services
    add_section_title(doc, "7. Services", "Services presents training offers, business services, and server-access products.")
    doc.add_paragraph("Browse services", style="Heading 2")
    add_bullets(doc, [
        "Search and open a service card to view its description, modules or skills, assigned trainers, price plans, and availability.",
        "Save a service to Bookmarks for later review.",
        "Open a trainer from the service page to review the trainer profile and ratings.",
        "Authorized staff can upload, edit, publish, or remove service listings according to their permission set.",
    ])
    doc.add_paragraph("Purchase server access", style="Heading 2")
    add_steps(doc, [
        "Open the required server-access service and select an available duration plan.",
        "Confirm the displayed amount and proceed to payment.",
        "Complete payment in the payment window.",
        "After payment verification, the assigned login details are delivered through the protected service chat.",
        "Open History to confirm payment status and Open login chat to retrieve the delivered access details.",
    ])
    add_callout(doc, "Availability", "A plan can show Out of stock when no server credential is available. Administrators add credentials and control the inventory shown to users.", "warning")
    doc.add_paragraph("Staff controls", style="Heading 2")
    add_bullets(doc, [
        "Upload a service with title, description, pricing, media, type, and trainer associations.",
        "For server access, maintain available credentials and the service's public visibility.",
        "Use payment history to investigate confirmation or delivery exceptions.",
    ])

    add_page_break(doc)

    # Meetings scheduling
    add_section_title(doc, "8. Meetings: Plan, Invite, and Join", "Meetings can be hosted inside the SJ platform or linked to an external provider such as Zoom, JioMeet, Google Meet, or Microsoft Teams.")
    doc.add_paragraph("Meetings page", style="Heading 2")
    add_bullets(doc, [
        "Use the calendar or Upcoming / Previous tabs to find meetings.",
        "Start begins an immediate internal meeting when the account has permission.",
        "Join accepts a meeting link or meeting ID and password.",
        "Plan opens the scheduling form for authorized staff.",
    ])
    doc.add_paragraph("Plan a meeting", style="Heading 2")
    add_steps(doc, [
        "Enter the meeting title.",
        "Choose SJ Internal or External Link.",
        "Choose Group to invite an existing group, or Individuals to search and select specific accounts.",
        "Set the start date, optional final end date, start time, end time, duration, and recurrence.",
        "For weekly or monthly recurrence, select the required weekdays or calendar dates.",
        "For an external meeting, select the provider and enter its link plus optional provider meeting ID and password.",
        "Choose Schedule and notify. Invited accounts receive the meeting details in chat and notifications.",
    ])
    add_callout(doc, "Link lifecycle", "Each internal meeting receives its own link, meeting ID, and password. The meeting remains usable only within its scheduled lifecycle and is unavailable after the final end date and time.", "success")
    doc.add_paragraph("Manage a meeting", style="Heading 2")
    add_bullets(doc, [
        "Expand a meeting card to copy the link or full invitation, review participant count, and view credentials.",
        "Authorized staff can add participants, remove participants from reporting, edit details, connect the meeting to Dashboard, or delete the meeting.",
        "Only invited users see restricted meeting records and their allocated training dashboards. Staff can view the meetings they manage.",
    ])

    add_page_break(doc)

    # Meeting room
    add_section_title(doc, "9. Internal Meeting Room", "An internal meeting opens a preview first, then the live room.")
    doc.add_paragraph("Before joining", style="Heading 2")
    add_steps(doc, [
        "Open the internal meeting link or choose Join from the meeting card.",
        "Confirm the meeting ID and password when requested.",
        "Review the camera preview and microphone state.",
        "Confirm or edit the displayed participant name, then enter the room.",
    ])
    doc.add_paragraph("Room controls", style="Heading 2")
    add_table(
        doc,
        ["Control", "What it does"],
        [
            ("Mic", "Mutes or unmutes the participant microphone."),
            ("Video", "Turns the participant camera stream on or off."),
            ("Share", "Shares a selected screen, window, or browser tab."),
            ("Record", "Records the meeting locally on the current device and downloads the file when stopped."),
            ("Raise", "Signals the host and participants that the user wants to speak."),
            ("People", "Opens the participant list."),
            ("Chat", "Opens meeting chat."),
            ("Leave", "Ends this participant's attendance session and exits the room."),
        ],
        [1700, 7660],
    )
    add_callout(doc, "Local recording", "The recording is created by the participant's device and downloaded there. It is not automatically uploaded to the platform or stored for other users. Long recordings require sufficient device memory and free storage.", "warning")
    add_text(doc, "For reliable camera and microphone shutdown, use the in-room controls and then Leave when finished. The attendance record closes when the participant leaves or the room session ends.")

    add_page_break(doc)

    # Internal vs external
    add_section_title(doc, "10. Internal and External Meeting Attendance", "Both meeting types support invitations and progress tracking, but they provide different attendance detail.")
    add_table(
        doc,
        ["Capability", "SJ Internal", "External Link"],
        [
            ("Meeting room", "Runs inside the SJ platform", "Opens the selected external provider"),
            ("Join evidence", "Recorded when the user enters the room", "Recorded when the user selects Join during the scheduled session"),
            ("Join and leave time", "Tracked", "Actual provider leave time is not available"),
            ("Duration", "Calculated from attendance sessions", "Shown as Not tracked"),
            ("Progress", "Counts eligible attended days", "Join action can count the scheduled day"),
            ("Media quality", "Depends on device, network, and participant load", "Handled by the external provider"),
        ],
        [2200, 3580, 3580],
        font_size=9.1,
    )
    add_callout(doc, "Important", "For external meetings, the platform cannot confirm how long a person stayed after the external provider opens. Use external links for scale or provider features, and interpret attendance as a recorded join action unless provider integration is added later.", "danger")
    doc.add_paragraph("Absent status", style="Heading 2")
    add_text(doc, "After a scheduled class has ended, an allocated trainer or participant with no qualifying attendance entry for that day is shown as Absent. Before the class ends, the system waits rather than marking the person absent early.")
    doc.add_paragraph("Recommended use", style="Heading 2")
    add_bullets(doc, [
        "Use SJ Internal when exact join, leave, duration, and in-app collaboration matter most.",
        "Use External Link when the selected provider is preferred for larger sessions or provider-specific controls.",
        "Always connect training meetings to Dashboard before classes begin so the correct trainer and participant list are tracked.",
    ])

    add_page_break(doc)

    # Dashboard
    add_section_title(doc, "11. Dashboard", "Dashboard converts live requirement and meeting activity into operational reports.")
    doc.add_paragraph("Training Report", style="Heading 2")
    add_steps(doc, [
        "Expand an already created meeting and choose Connect to dashboard.",
        "Set total training days and training weekdays.",
        "Select the trainer and the invited participants to track. Staff may add or remove eligible accounts later.",
        "Open Dashboard, choose Training Report, and select a training card.",
        "Review course completion, planned and completed days, remaining days, tracked attendees, the attendance register, trainer progress, and participant percentages.",
    ])
    add_bullets(doc, [
        "Admin and authorized Employee accounts can review all training reports available to their role.",
        "Trainers and participants see only reports for meetings allocated to their account.",
        "Internal attendance combines join, rejoin, and leave activity by account and date.",
        "The selected report refreshes automatically while the Training Report tab is open.",
    ])
    doc.add_paragraph("Employee Report", style="Heading 2")
    add_bullets(doc, [
        "Requirements shows the total internal requirement records.",
        "Active shows work still open or in sourcing.",
        "Profiles added totals candidate profiles across requirements.",
        "Completed shows closed requirements.",
        "Each row connects the requirement creator to sourcing contributors, profile counts, time spent, and the employee who closed the requirement.",
        "View details opens the matching Task Board record.",
    ])
    add_callout(doc, "Live reporting", "Attendance and task data refresh automatically. A short delay is normal because the report consolidates activity from multiple accounts and devices.", "success")

    add_page_break(doc)

    # Trainers users and bookmarks
    add_section_title(doc, "12. Trainers / Users and Bookmarks", "These pages help people find accounts, review trainer credentials, and keep useful content for later.")
    doc.add_paragraph("Trainers / Users", style="Heading 2")
    add_bullets(doc, [
        "Search by name, role, or specialization.",
        "Open a profile to view the account's role, professional details, experience, location, training mode, biography, and available resume information.",
        "The green presence indicator shows that the account is currently active in the app. It may take a short moment to clear after a device closes or loses connection.",
        "Eligible users can rate or review a trainer. A person cannot review their own profile.",
        "Use Chat from a profile when direct contact is permitted. Otherwise, the app may offer a request to Admin Service.",
    ])
    doc.add_paragraph("Bookmarks", style="Heading 2")
    add_bullets(doc, [
        "Saved Services contains service cards bookmarked by the current account.",
        "Saved Posts contains feed posts bookmarked by the current account.",
        "Open an item to return to its details, or remove the bookmark when it is no longer needed.",
        "Bookmarks are personal to the signed-in account.",
    ])

    add_page_break(doc)

    # History and Requests
    add_section_title(doc, "13. History and Requests", "History tracks server-access payments. Requests controls approval-based chat access.")
    doc.add_paragraph("Payment History", style="Heading 2")
    add_bullets(doc, [
        "A normal user sees payments made by their own account.",
        "Admin and Super Admin can view payment activity across all accounts.",
        "Status filters include All, Paid, Pending, Failed, and Cancelled.",
        "Admin summary cards show paid, pending, failed, cancelled, and paid total values.",
        "Each record shows the service, customer where permitted, plan length, amount, order ID, payment ID, and status.",
        "Open login chat opens the conversation containing delivered credentials.",
        "Finish delivery retries confirmation when payment evidence exists but access delivery is incomplete.",
    ])
    add_callout(doc, "Live payment view", "Payment History checks for updates automatically every five seconds. Use Refresh for an immediate check after investigating a payment.", "success")
    doc.add_paragraph("Chat Requests", style="Heading 2")
    add_steps(doc, [
        "When a direct chat is not yet permitted, the requester chooses to send an access request to Admin Service.",
        "Authorized staff open Requests to review the requester, requested contact, date, and current status.",
        "Choose Approve to create or enable the conversation, or Reject to decline it.",
        "The requester receives the updated result and can chat after approval.",
    ])

    add_page_break(doc)

    # Account management
    add_section_title(doc, "14. Account Management", "Account Management is available to Admin and Super Admin accounts and is the main place to create and control people and teams.")
    doc.add_paragraph("Find and review accounts", style="Heading 2")
    add_bullets(doc, [
        "Use Dashboard, Users, Trainers, and Employees tabs to switch account groups.",
        "Search by name, email, role, specialization, or team.",
        "Open Details or Activity to review an account.",
        "Use View As to confirm what that account can see. Exit impersonation before making administrative changes.",
    ])
    doc.add_paragraph("Create an employee", style="Heading 2")
    add_steps(doc, [
        "Create or select the employee's team, such as Business Development, Sourcing, HR, or Accounting.",
        "Enter the employee name, email, and temporary password.",
        "Assign the team.",
        "Select only the permissions needed for the role, such as viewing users or chats, posting feeds or services, arranging meetings, managing chat requests, or all access.",
        "Create the account and securely share its sign-in details with the employee.",
    ])
    doc.add_paragraph("Administrative actions", style="Heading 2")
    add_bullets(doc, [
        "Restrict temporarily removes privileged access without deleting the account.",
        "Unrestrict restores permitted access.",
        "Delete permanently removes the selected account and should be used only after confirming retention needs.",
        "Teams can be added for new business functions or removed when no longer used.",
    ])
    add_callout(doc, "Control principle", "Use individual accounts and minimum required permissions. This keeps chat, attendance, payment, sourcing, and administrative activity correctly attributed.", "warning")

    add_page_break(doc)

    # Data Management and Settings
    add_section_title(doc, "15. Data Management and Settings", "Data Management controls stored records. Settings controls the current user's profile, security, notifications, and media preferences.")
    doc.add_paragraph("Data Management - administrators only", style="Heading 2")
    add_bullets(doc, [
        "Stored Media lists files used in chat or feed locations and supports preview and selection.",
        "Delete cloud file removes the stored file. Existing messages or posts may then show that the attachment is unavailable.",
        "Delete selected removes multiple chosen media files.",
        "Message Records lists stored messages and provides record deletion controls.",
    ])
    add_callout(doc, "Permanent action", "Deletion in Data Management can affect what users can open later. Confirm business retention requirements before removing files or records.", "danger")
    doc.add_paragraph("Settings", style="Heading 2")
    add_table(
        doc,
        ["Area", "What the user can manage"],
        [
            ("Profile", "Photo, name, professional title, and role-specific trainer details"),
            ("Password", "Change the current account password"),
            ("Notifications", "Push, course, message, and announcement preferences"),
            ("Chat and media", "Automatic media download and local media behavior"),
            ("Delete account", "Permanently remove the current account after password confirmation"),
        ],
        [2100, 7260],
    )
    add_callout(doc, "Push setup", "Choose Enable notifications in Settings when permission has not yet been granted. If permission was previously blocked, it must be restored from the browser or device site settings.", "info")

    add_page_break(doc)

    # Realtime explanation
    add_section_title(doc, "16. How Live Updates Work", "The platform combines immediate on-screen feedback, regular background synchronization, meeting-room activity, and device push notifications.")
    doc.add_paragraph("What users experience", style="Heading 2")
    add_table(
        doc,
        ["Activity", "Expected live behavior"],
        [
            ("Chat and presence", "Messages, unread counts, delivery status, seen status, and online indicators update automatically, normally within a few seconds."),
            ("Feed and requirements", "New posts, likes, comments, and requirement changes appear without refreshing while the app is active."),
            ("Task Board", "Worker joins, candidate profiles, mentions, and status decisions synchronize to connected staff."),
            ("Meeting room", "Audio, video, screen sharing, room chat, and participant state operate live during the session."),
            ("Training Dashboard", "Attendance and progress refresh while the report is open; join, rejoin, leave, and absent results are consolidated."),
            ("Payments", "History refreshes automatically every five seconds after payment or delivery changes."),
            ("Scheduled work", "Scheduled messages and reminders are checked by the platform scheduler and may arrive within the scheduler's check interval rather than at the exact second."),
            ("Push notifications", "The device receives eligible alerts outside the open app when permission, internet, and device notification settings allow it."),
        ],
        [2300, 7060],
        font_size=9.1,
    )
    doc.add_paragraph("What happens when a device is offline", style="Heading 2")
    add_bullets(doc, [
        "Actions that require the internet may wait, fail, or show an unsent state.",
        "The latest feed, chat, notification, task, and account state loads after the connection returns.",
        "An online indicator may remain visible briefly while the previous presence state expires.",
        "Push delivery can be delayed by battery saving, background restrictions, Do Not Disturb, or browser policy.",
    ])
    add_callout(doc, "Meaning of real time", "In this guide, real time means automatic synchronization without a manual page refresh. It does not guarantee zero-delay delivery; internet quality, device state, and the action type can introduce a short delay.", "info")

    add_page_break(doc)

    # Operating checklist/troubleshooting
    add_section_title(doc, "17. Daily Operating Checklist", "A short routine keeps the platform accurate and dependable for the whole organization.")
    doc.add_paragraph("Administrator / employee start of day", style="Heading 2")
    add_bullets(doc, [
        "Review Notifications and Requests.",
        "Check today's meetings, invited accounts, and meeting passwords.",
        "Confirm training meetings are connected to Dashboard with the correct trainer and participants.",
        "Review active requirements and candidate profile decisions on Task Board.",
        "Review pending or failed payments and server credential availability.",
    ])
    doc.add_paragraph("Before a training session", style="Heading 2")
    add_bullets(doc, [
        "Confirm the date, time, recurrence, meeting type, and end date.",
        "Confirm the group or selected individuals and add any late participant.",
        "For internal meetings, test microphone, camera, screen share, and recording on the actual device.",
        "For external meetings, test the provider link and remember that actual duration is not tracked by the SJ platform.",
    ])
    doc.add_paragraph("End of day", style="Heading 2")
    add_bullets(doc, [
        "Review absent records only after scheduled sessions have ended.",
        "Update candidate profiles to Follow up, Rejected, or Completed where a decision was made.",
        "Close requirements only when the work is fully finished because closure is permanent.",
        "Log out from shared devices.",
    ])
    doc.add_paragraph("Quick troubleshooting", style="Heading 2")
    add_table(
        doc,
        ["Issue", "First check"],
        [
            ("No notification", "Permission, device volume, notification channel, Do Not Disturb, and internet"),
            ("Page not updating", "Internet, app foreground state, then reopen the page or app"),
            ("Cannot see a page", "Account role, restriction state, and employee permissions"),
            ("Cannot see a meeting/report", "Whether the account was invited or added to the dashboard report"),
            ("Camera or mic issue", "Browser permission, selected device, mute/video state, and another app using the device"),
            ("External duration missing", "Expected behavior; only the external join action is recorded"),
            ("Payment pending", "Wait for automatic refresh, use Refresh, then use Finish delivery when eligible"),
        ],
        [2600, 6760],
        font_size=9.2,
    )
    add_callout(doc, "Support handover", "When reporting an issue, include the account name, role, page, action attempted, approximate time, device type, and a screenshot. Do not send passwords in screenshots or chat.", "warning")

    # Final page
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(1.8))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("SJ INFO BUSINESS SOLUTIONS")
    set_run_font(r, size=21, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("One platform for communication, services, meetings, training progress, and internal operations.")
    set_run_font(r, size=12, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("End of client user guide")
    set_run_font(r, size=10, color=GREEN, bold=True)

    # Core document properties
    props = doc.core_properties
    props.title = "SJ Info Business Solutions Platform User Guide"
    props.subject = "Client-facing guide to platform pages, workflows, and live behavior"
    props.author = "SJ Info Business Solutions"
    props.keywords = "SJ Info Business Solutions, user guide, platform, meetings, attendance, task board"
    props.comments = "Prepared for client onboarding and operational handover."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
